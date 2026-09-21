"""Research-only multi-week stability matrix using the existing author-replica MT5 implementation.

Weeks are Monday-Friday UTC. No canonical Strategy A rule is defined here.
Each week is evaluated independently with the same fixed configuration.
"""
from __future__ import annotations
import importlib.util, json, os
from collections import defaultdict
from datetime import datetime, timezone, timedelta
from pathlib import Path
import MetaTrader5 as mt5
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

SOURCE = Path("scripts/run-author-replica-mt5-nonoverlap-stability.py")
SYMBOL = os.getenv("TRADING_SYMBOL", "XAUUSD.ecn")
N = int(os.getenv("BARS", "10000"))
spec = importlib.util.spec_from_file_location("sp2l_replica", SOURCE)
mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)

# Four complete Mon-Fri UTC weeks ending 2026-09-18, including the week just tested.
WEEKS = [
    ("2026-09-14", "2026-09-18"),
    ("2026-09-07", "2026-09-11"),
    ("2026-08-31", "2026-09-04"),
    ("2026-08-24", "2026-08-28"),
]
OUT = Path("artifacts/SP2L_author_replica_4week_stability_matrix.xlsx")
JSON_OUT = OUT.with_suffix(".json")

def dt(s, end=False):
    return datetime.fromisoformat(s + ("T23:59:59+00:00" if end else "T00:00:00+00:00"))

def U(ts):
    return datetime.fromtimestamp(int(ts), timezone.utc).isoformat()

def run_week(start, end):
    # End-anchored lookback preserves the existing MT5 acquisition pattern.
    rates = mt5.copy_rates_from(SYMBOL, mt5.TIMEFRAME_M1, end, N)
    if rates is None or len(rates) == 0:
        return {"week_start_utc": start.isoformat(), "week_end_utc": end.isoformat(),
                "error": str(mt5.last_error())}
    c = [{"time": int(r[0]), "open": float(r[1]), "high": float(r[2]), "low": float(r[3]), "close": float(r[4])} for r in rates]
    rows = []
    for i in range(4, len(c)):
        s = mod.signal(c, i)
        if not s: continue
        direction, entry, sl = s
        signal_ts = c[i-1]["time"]
        if not (start.timestamp() <= signal_ts <= end.timestamp()): continue
        risk = abs(entry-sl)
        if risk <= 0 or risk > mod.MAX_SL: continue
        tp = entry + mod.TP_R*risk if direction == "BUY" else entry - mod.TP_R*risk
        outcome, reason, exit_ts = "OPEN_OR_UNRESOLVED", "NO_SL_TP_HIT_IN_AVAILABLE_DATA", None
        for j in range(i+1, len(c)):
            x = c[j]
            hit_sl = x["low"] <= sl if direction == "BUY" else x["high"] >= sl
            hit_tp = x["high"] >= tp if direction == "BUY" else x["low"] <= tp
            if hit_sl and hit_tp:
                outcome, reason, exit_ts = "AMBIGUOUS", "SL_AND_TP_SAME_BAR", x["time"]; break
            if hit_sl:
                outcome, reason, exit_ts = "LOSS", "SL", x["time"]; break
            if hit_tp:
                outcome, reason, exit_ts = "WIN", "TP", x["time"]; break
        rows.append({"time": U(signal_ts), "date": U(signal_ts)[:10], "direction": direction,
                     "entry": entry, "sl": sl, "tp": tp, "risk": risk,
                     "outcome": outcome, "R": (1 if outcome=="WIN" else -1 if outcome=="LOSS" else None),
                     "exit_time": U(exit_ts) if exit_ts else "", "reason": reason})
    w = sum(x["outcome"]=="WIN" for x in rows); l = sum(x["outcome"]=="LOSS" for x in rows)
    a = sum(x["outcome"]=="AMBIGUOUS" for x in rows); u = sum(x["outcome"]=="OPEN_OR_UNRESOLVED" for x in rows)
    b = sum(x["direction"]=="BUY" for x in rows); bw = sum(x["direction"]=="BUY" and x["outcome"]=="WIN" for x in rows)
    bl = sum(x["direction"]=="BUY" and x["outcome"]=="LOSS" for x in rows)
    s = sum(x["direction"]=="SELL" for x in rows); sw = sum(x["direction"]=="SELL" and x["outcome"]=="WIN" for x in rows)
    sl = sum(x["direction"]=="SELL" and x["outcome"]=="LOSS" for x in rows)
    dec=w+l
    equity=peak=dd=0
    for x in rows:
        if x["R"] is not None:
            equity += x["R"]; peak=max(peak,equity); dd=max(dd,peak-equity)
    return {"week_start_utc":start.isoformat(),"week_end_utc":end.isoformat(),
            "returned_bars":len(c),"first_returned_utc":U(c[0]["time"]),"last_returned_utc":U(c[-1]["time"]),
            "signals":len(rows),"wins":w,"losses":l,"ambiguous":a,"open_or_unresolved":u,
            "decisive_win_rate":w/dec if dec else None,"total_R":w-l,"profit_factor":w/l if l else None,
            "max_drawdown_R":dd,"buy_signals":b,"buy_wins":bw,"buy_losses":bl,
            "buy_decisive_wr":bw/(bw+bl) if bw+bl else None,"sell_signals":s,"sell_wins":sw,"sell_losses":sl,
            "sell_decisive_wr":sw/(sw+sl) if sw+sl else None,"trades":rows}

if not mt5.initialize(): raise SystemExit(f"MT5 initialize failed: {mt5.last_error()}")
try:
    if not mt5.symbol_select(SYMBOL, True): raise SystemExit(f"symbol_select failed: {mt5.last_error()}")
    results=[run_week(dt(a),dt(b,True)) for a,b in WEEKS]
finally:
    mt5.shutdown()

# Aggregate only decisive outcomes.
tw=sum(x.get("wins",0) for x in results); tl=sum(x.get("losses",0) for x in results)
agg={"weeks":len(results),"wins":tw,"losses":tl,"ambiguous":sum(x.get("ambiguous",0) for x in results),
     "open_or_unresolved":sum(x.get("open_or_unresolved",0) for x in results),
     "decisive_win_rate":tw/(tw+tl) if tw+tl else None,"total_R":tw-tl,"profit_factor":tw/tl if tl else None}

wb=Workbook(); ws=wb.active; ws.title="Stability Matrix"
headers=["week_start_utc","week_end_utc","signals","wins","losses","ambiguous","open_or_unresolved","decisive_win_rate","total_R","profit_factor","max_drawdown_R","buy_signals","buy_wins","buy_losses","buy_decisive_wr","sell_signals","sell_wins","sell_losses","sell_decisive_wr","first_returned_utc","last_returned_utc"]
ws.append(headers)
for x in results: ws.append([x.get(h) for h in headers])
ws2=wb.create_sheet("Aggregate"); ws2.append(["metric","value"])
for k,v in agg.items(): ws2.append([k,v])
ws3=wb.create_sheet("All Signals"); sh=["week_start_utc","time","date","direction","entry","sl","tp","risk","outcome","R","exit_time","reason"]; ws3.append(sh)
for x in results:
    for t in x.get("trades",[]): ws3.append([x["week_start_utc"]]+[t.get(h) for h in sh[1:]])
for sheet in wb.worksheets:
    sheet.freeze_panes="A2"
    for col in sheet.columns:
        sheet.column_dimensions[get_column_letter(col[0].column)].width=min(max(len(str(c.value or "")) for c in col)+2,40)
OUT.parent.mkdir(exist_ok=True); wb.save(OUT)
JSON_OUT.write_text(json.dumps({"research_only":True,"symbol":SYMBOL,"timeframe":"M1","requested_bars":N,"config":{"pGapPrice":mod.P_GAP,"spikeMultiplier":mod.SPIKE_MULT,"maxSlPrice":mod.MAX_SL,"tpR":mod.TP_R},"weeks":results,"aggregate":agg},indent=2),encoding="utf-8")
print(json.dumps({"xlsx":str(OUT),"json":str(JSON_OUT),"aggregate":agg,"weeks":[{k:v for k,v in x.items() if k!="trades"} for x in results]},indent=2))
