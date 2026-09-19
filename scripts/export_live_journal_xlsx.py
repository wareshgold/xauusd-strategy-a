"""Export the live journal into a single analysis workbook.

Requires openpyxl in the local MT5 Python environment.
Raw JSONL remains the audit source; XLSX is a derived analysis view.
"""

from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from live_journal import read_jsonl, SIGNALS, TRADES, SNAPSHOTS

OUT = Path("runtime/exports/SP2L_Live_Trade_Journal.xlsx")


def write_sheet(wb: Workbook, name: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(name)
    if not rows:
        ws.append(["No records"])
        return

    keys: list[str] = []
    seen: set[str] = set()
    for row in rows:
        for key in row:
            if key not in seen:
                seen.add(key)
                keys.append(key)

    ws.append(keys)
    for row in rows:
        ws.append([row.get(key) for key in keys])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, key in enumerate(keys, start=1):
        width = max(12, min(42, len(key) + 2))
        ws.column_dimensions[get_column_letter(index)].width = width


def write_summary(wb: Workbook, signals: list[dict], trades: list[dict]) -> None:
    ws = wb.create_sheet("Summary", 0)
    closed = [r for r in trades if str(r.get("status", "")).upper() == "CLOSED"]
    wins = [r for r in closed if str(r.get("result", "")).upper() == "WIN"]
    losses = [r for r in closed if str(r.get("result", "")).upper() == "LOSS"]
    r_values = [float(r["r_multiple"]) for r in closed if r.get("r_multiple") is not None]
    positive = sum(v for v in r_values if v > 0)
    negative = -sum(v for v in r_values if v < 0)

    rows = [
        ("Generated UTC", __import__("datetime").datetime.now(__import__("datetime").timezone.utc).isoformat()),
        ("Signal records", len(signals)),
        ("Trade records", len(trades)),
        ("Closed trades", len(closed)),
        ("Wins", len(wins)),
        ("Losses", len(losses)),
        ("Win rate %", 100 * len(wins) / (len(wins) + len(losses)) if wins or losses else None),
        ("Sum R", sum(r_values) if r_values else 0),
        ("Profit factor by R", positive / negative if negative else None),
    ]
    ws.append(["Metric", "Value"])
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 30


def main() -> None:
    signals = read_jsonl(SIGNALS)
    trades = read_jsonl(TRADES)
    snapshots = read_jsonl(SNAPSHOTS)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    write_summary(wb, signals, trades)
    write_sheet(wb, "Signals", signals)
    write_sheet(wb, "Trades", trades)
    write_sheet(wb, "Market Snapshots", snapshots)
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
